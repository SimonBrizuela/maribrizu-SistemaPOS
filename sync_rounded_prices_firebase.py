"""
Sube a Firebase los precios redondeados que ya están en SQLite.
"""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))

import sqlite3
from datetime import datetime, timezone
from pos_system.config import DATABASE_PATH

def main():
    conn = sqlite3.connect(str(DATABASE_PATH))
    cur = conn.cursor()

    # Traer todos los productos con precio multiplo de 100 que pueden haber cambiado
    # (identificamos los que matchean los nombres del Excel)
    import openpyxl
    wb = openpyxl.load_workbook('productos_precio_irregular.xlsx')
    ws = wb.active
    nombres_excel = set()
    barcodes_excel = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        nombre, rubro, cat, precio, barcode, stock = row
        if nombre:
            nombres_excel.add(nombre)
        if barcode:
            barcodes_excel.add(str(barcode))

    print(f"Productos en Excel: {len(nombres_excel)} nombres, {len(barcodes_excel)} barcodes")

    # Obtener esos productos de SQLite (ya con precio redondeado)
    placeholders_n = ','.join('?' * len(nombres_excel))
    cur.execute(
        f'SELECT id, name, category, price, cost, stock, discount_value FROM products WHERE name IN ({placeholders_n})',
        list(nombres_excel)
    )
    products = cur.fetchall()
    print(f"Encontrados en SQLite: {len(products)} productos")

    # Conectar Firebase
    print("Conectando a Firebase...")
    import firebase_admin
    from firebase_admin import credentials, firestore

    try:
        firebase_admin.get_app()
        db_fb = firestore.client()
    except ValueError:
        cred = credentials.Certificate('firebase_key.json')
        firebase_admin.initialize_app(cred)
        db_fb = firestore.client()

    print("Firebase conectado. Subiendo en batches...")

    BATCH_SIZE = 500
    total = 0
    col = db_fb.collection('inventario')
    ts = datetime.now(timezone.utc)

    for i in range(0, len(products), BATCH_SIZE):
        chunk = products[i:i+BATCH_SIZE]
        batch = db_fb.batch()
        for db_id, name, category, price, cost, stock_val, discount in chunk:
            ref = col.document(str(db_id))
            batch.set(ref, {
                'id':                   db_id,
                'nombre':               name or '',
                'categoria':            category or 'Sin categoría',
                'precio':               float(price or 0),
                'costo':                float(cost or 0),
                'stock':                int(stock_val or 0),
                'descuento':            float(discount or 0),
                'ultima_actualizacion': ts,
            }, merge=True)
        batch.commit()
        total += len(chunk)
        print(f"  Batch {i//BATCH_SIZE + 1}: {len(chunk)} productos subidos (total: {total})")

    print(f"\n✓ {total} productos con precios redondeados subidos a Firebase.")
    conn.close()

if __name__ == '__main__':
    main()
